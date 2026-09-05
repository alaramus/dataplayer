"""Reading a spreadsheet workbook: its sheets, and a bounded page of one.

A workbook is a container, and each of its worksheets is a node (D51). Nothing
here reads more than it was asked for: the zip's own directory is checked
before a byte is inflated, and rows are streamed through openpyxl's read-only
iterator, which stops one row past the page.
"""

import zipfile
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from app.adapters.cells import encode_cell
from app.adapters.delimited import fit, usable_names
from app.adapters.fetch import FetchColumn, Filter
from app.adapters.file.readers import Page, matcher, select
from app.adapters.relation import ColumnType

# A workbook is a zip, so its size on disk says nothing about what opening it
# costs. Both ceilings are read from metadata: the file's own size, and the
# uncompressed size the central directory declares for every entry.
MAX_WORKBOOK_BYTES = 64 * 1024 * 1024
MAX_UNPACKED_BYTES = 512 * 1024 * 1024
MAX_ENTRIES = 10_000

# Excel allows 16,384 columns in a sheet, and a page of a thousand rows across
# all of them is bounded by the format and by nothing else.
MAX_COLUMNS = 1024

# A worksheet column has no type: Excel types cells, so a type read off one
# page would change with the page, exactly as it would for a CSV.
TEXT = ColumnType(sql="text", category="string")

# Per-value types, for the canonical rendering shared with every other source.
# The choice is per cell because that is how the format itself works.
_BOOLEAN = ColumnType(sql="boolean", category="boolean")
_DATETIME = ColumnType(sql="timestamp", category="datetime")
_NUMBER = ColumnType(sql="numeric", category="number")
_BINARY = ColumnType(sql="bytea", category="binary")


@dataclass(frozen=True)
class SheetFacts:
    """One worksheet as the listing of a workbook describes it."""

    position: int
    title: str
    hidden: bool
    # From the dimension the sheet declares, header row excluded. A workbook
    # written without one states neither.
    rows: int | None
    columns: int | None


class NotAWorkbook(ValueError):
    """The bytes under a workbook's name are not one."""


def sheets(path: Path) -> list[SheetFacts]:
    """Every worksheet of a workbook, in workbook order. Blocking."""
    with _opened(path) as workbook:
        return [
            _facts(position, sheet)
            for position, sheet in enumerate(workbook.worksheets, start=1)
        ]


def read_page(
    path: Path,
    position: int,
    limit: int,
    offset: int,
    filters: Sequence[Filter] = (),
) -> Page:
    """One bounded page of one worksheet. Blocking.

    Skipping to an offset walks the rows before it: XLSX has no row index, so
    the tenth page costs ten pages of iteration, and a filter costs the same
    walk over the whole sheet (D109).
    """
    with _opened(path) as workbook:
        sheet = _worksheet(workbook, position)
        rows = sheet.iter_rows(values_only=True)
        names = _header(rows)
        if names is None:
            if filters:
                raise ValueError("This sheet has no columns to filter by")
            return Page(columns=[], rows=[], has_more=False, positions=[])

        match = matcher(filters, names)
        page, positions, has_more = select(
            (
                fit([_cell_text(value) for value in raw], len(names))
                for raw in rows
            ),
            limit,
            offset,
            match,
        )

    return Page(
        columns=[FetchColumn(name=name, type=TEXT) for name in names],
        rows=page,
        has_more=has_more,
        positions=positions,
    )


def read_header(path: Path, position: int) -> tuple[list[str], SheetFacts, int]:
    """The column names of one worksheet, its facts, and how many there are.

    All three from one open: the relation model needs every one of them, and a
    workbook opened twice is a workbook read twice. Blocking.
    """
    with _opened(path) as workbook:
        sheet = _worksheet(workbook, position)
        names = _header(sheet.iter_rows(values_only=True))
        return names or [], _facts(position, sheet), len(workbook.worksheets)


def _facts(position: int, sheet: Any) -> SheetFacts:
    """One worksheet described from what it declares about itself.

    `sheet` is typed loosely on purpose: openpyxl hands back a read-only
    worksheet whose stubs share no protocol with the writable one, and every
    attribute read here is converted before it leaves.
    """
    max_row = sheet.max_row
    max_column = sheet.max_column
    return SheetFacts(
        position=position,
        title=str(sheet.title),
        hidden=str(sheet.sheet_state) != "visible",
        rows=max(int(max_row) - 1, 0) if max_row is not None else None,
        columns=int(max_column) if max_column is not None else None,
    )


def _header(rows: Iterator[Any]) -> list[str] | None:
    """Column names from the first row, or None for an empty sheet."""
    first = next(rows, None)
    if first is None:
        return None
    if len(first) > MAX_COLUMNS:
        raise ValueError(
            f"This sheet has more than {MAX_COLUMNS} columns, "
            "which is more than one page may carry"
        )
    return usable_names([_header_text(value) for value in first])


def _header_text(value: Any) -> str:
    """A heading cell as a name. An empty one is named by its position later."""
    text = _cell_text(value)
    return "" if text is None else text


def _cell_text(value: Any) -> str | None:
    """One cell as text, through the rendering every other source uses.

    The type is chosen per value rather than per column because a spreadsheet
    types cells: one column may hold a date, a number and a word.
    """
    if isinstance(value, bool):
        return encode_cell(value, _BOOLEAN)
    if isinstance(value, datetime | date | time):
        return encode_cell(value, _DATETIME)
    if isinstance(value, int | float | Decimal):
        return encode_cell(value, _NUMBER)
    if isinstance(value, bytes | bytearray):
        return encode_cell(value, _BINARY)
    if isinstance(value, timedelta):
        # Excel keeps a duration as a number of days; openpyxl hands it back as
        # a timedelta, and its own spelling is the only lossless one here.
        return str(value)
    return encode_cell(value, TEXT)


def _worksheet(workbook: Workbook, position: int) -> Any:
    """One worksheet by its one-based position among the worksheets.

    Loosely typed for the reason `_facts` is: a read-only worksheet and a
    writable one share no type in openpyxl's stubs.
    """
    worksheets = workbook.worksheets
    if position < 1 or position > len(worksheets):
        raise ValueError(f"This workbook has no sheet {position}")
    return worksheets[position - 1]


class _opened:
    """The workbook, guarded before it is opened and closed after it is read.

    A context manager rather than a function because a read-only workbook holds
    the file open until it is closed, and every reader here must close it.
    """

    _workbook: Workbook

    def __init__(self, path: Path) -> None:
        self._path = path

    def __enter__(self) -> Workbook:
        _guard(self._path)
        try:
            self._workbook = load_workbook(
                self._path, read_only=True, data_only=True
            )
        except zipfile.BadZipFile:
            raise NotAWorkbook("This file is not a workbook") from None
        except (KeyError, ValueError, TypeError) as exc:
            # openpyxl raises whatever the malformed part of the package makes
            # it raise. The reason is ours, and never the library's text.
            raise NotAWorkbook("This workbook could not be read") from exc
        return self._workbook

    def __exit__(self, *_: object) -> None:
        self._workbook.close()


def _guard(path: Path) -> None:
    """Refuse a workbook that would cost too much, before it is opened.

    Every number checked here is metadata: the size of the file, and the sizes
    the zip's central directory declares. Nothing is decompressed to find out.
    """
    size = path.stat().st_size
    if size > MAX_WORKBOOK_BYTES:
        raise ValueError(
            f"This workbook is larger than {MAX_WORKBOOK_BYTES // (1024 * 1024)} MB"
        )
    try:
        with zipfile.ZipFile(path) as archive:
            entries: Sequence[zipfile.ZipInfo] = archive.infolist()
    except zipfile.BadZipFile:
        raise NotAWorkbook("This file is not a workbook") from None

    if len(entries) > MAX_ENTRIES:
        raise ValueError(f"This workbook holds more than {MAX_ENTRIES} parts")
    declared = sum(entry.file_size for entry in entries)
    if declared > MAX_UNPACKED_BYTES:
        raise ValueError(
            "This workbook unpacks to more than "
            f"{MAX_UNPACKED_BYTES // (1024 * 1024)} MB"
        )
