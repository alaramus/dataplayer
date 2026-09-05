"""Put sample files into a user's personal directory, and tidy up after tests.

A development tool, not part of the API: a fresh directory is empty, so there
would be nothing to look at or to run an end-to-end scenario against. It also
removes the directories of users that no longer exist, which is what the test
suite leaves behind every time it empties the shared development database.

    python -m app.seed_storage [email]
"""

import asyncio
import base64
import json
import logging
import math
import struct
import sys
import uuid
import wave
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from app.core.config import settings
from app.core.db import async_session_maker, loop_factory
from app.users import crud, storage

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# More than one page of the default size, so paging is exercised.
SAMPLE_ROWS = 250

SAMPLE_RATE = 8000
SAMPLE_ARCHIVE_README = """# Inside an archive

This document lives in `bundle.zip` and is read without unpacking it.
The folder beside it holds a text file, a table and a picture.
"""
SAMPLE_PDF_PAGES = [
    [
        "DataPlayer",
        "",
        "This document exists so the PDF viewer has something to draw.",
        "It is rendered by the application rather than by the browser's",
        "own viewer, so nothing in it can run.",
    ],
    [
        "Page two",
        "",
        "Pages are drawn as they are scrolled to.",
    ],
]

# A PostgreSQL log, which is the shape a server this application starts writes.
SAMPLE_LOG_LINES = [
    (1, "LOG", "database system is ready to accept connections", ""),
    (2, "LOG", "checkpoint starting: time", ""),
    (3, "ERROR", 'relation "orders" does not exist at character 15', "  STATEMENT:  SELECT * FROM orders\n"),
    (4, "WARNING", "there is already a transaction in progress", ""),
    (5, "LOG", "duration: 1201.884 ms  execute <unnamed>", "  DETAIL:  parameters: $1 = '42'\n"),
    (6, "FATAL", 'password authentication failed for user "someone"', ""),
    (7, "LOG", "received smart shutdown request", ""),
]

# One real 1×1 PNG: the viewer verifies bytes, so a signature followed by
# sample text would advertise an image that correctly refuses to open.
SAMPLE_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)

# Two seconds of real H.264 in an MP4, for the same reason the PNG is real: a
# hand-built container verifies and then fails in the player, which is a worse
# sample than none. 1.8 KB, encoded here because nothing binary is committed.
SAMPLE_MP4 = base64.b64decode(
    "".join(
        [
    "AAAAIGZ0eXBpc29tAAACAGlzb21pc28yYXZjMW1wNDEAAANkbW9vdgAAAGxtdmhkAAAAAAAA"
    "AAAAAAAAAAAD6AAAB9AAAQAAAQAAAAAAAAAAAAAAAAEAAAAAAAAAAAAAAAAAAAABAAAAAAAA"
    "AAAAAAAAAABAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAgAAAo50cmFrAAAAXHRr"
    "aGQAAAADAAAAAAAAAAAAAAABAAAAAAAAB9AAAAAAAAAAAAAAAAAAAAAAAAEAAAAAAAAAAAAA"
    "AAAAAAABAAAAAAAAAAAAAAAAAABAAAAAAKAAAAB4AAAAAAAkZWR0cwAAABxlbHN0AAAAAAAA"
    "AAEAAAfQAAAAAAABAAAAAAIGbWRpYQAAACBtZGhkAAAAAAAAAAAAAAAAAABAAAAAgABVxAAA"
    "AAAALWhkbHIAAAAAAAAAAHZpZGUAAAAAAAAAAAAAAABWaWRlb0hhbmRsZXIAAAABsW1pbmYA"
    "AAAUdm1oZAAAAAEAAAAAAAAAAAAAACRkaW5mAAAAHGRyZWYAAAAAAAAAAQAAAAx1cmwgAAAA"
    "AQAAAXFzdGJsAAAAuXN0c2QAAAAAAAAAAQAAAKlhdmMxAAAAAAAAAAEAAAAAAAAAAAAAAAAA"
    "AAAAAKAAeABIAAAASAAAAAAAAAABFUxhdmM2Mi4yOC4xMDIgbGlieDI2NAAAAAAAAAAAAAAA"
    "GP//AAAAL2F2Y0MBQsAe/+EAF2dCwB7ZAoR+XARAAAADAEAAAAQDxYuSAQAFaMuDyyAAAAAQ"
    "cGFzcAAAAAEAAAABAAAAFGJ0cnQAAAAAAAANrAAAAAAAAAAYc3R0cwAAAAAAAAABAAAAEAAA"
    "CAAAAAAUc3RzcwAAAAAAAAABAAAAAQAAABxzdHNjAAAAAAAAAAEAAAABAAAAEAAAAAEAAABU"
    "c3RzegAAAAAAAAAAAAAAEAAAAskAAAALAAAACwAAAAsAAAALAAAACwAAAAsAAAALAAAACwAA"
    "AAsAAAALAAAACwAAAAsAAAAKAAAACgAAAAoAAAAUc3RjbwAAAAAAAAABAAADlAAAAGJ1ZHRh"
    "AAAAWm1ldGEAAAAAAAAAIWhkbHIAAAAAAAAAAG1kaXJhcHBsAAAAAAAAAAAAAAAALWlsc3QA"
    "AAAlqXRvbwAAAB1kYXRhAAAAAQAAAABMYXZmNjIuMTIuMTAyAAAACGZyZWUAAANzbWRhdAAA"
    "AnAGBf//bNxF6b3m2Ui3lizYINkj7u94MjY0IC0gY29yZSAxNjUgcjMyMjMgMDQ4MGNiMCAt"
    "IEguMjY0L01QRUctNCBBVkMgY29kZWMgLSBDb3B5bGVmdCAyMDAzLTIwMjUgLSBodHRwOi8v"
    "d3d3LnZpZGVvbGFuLm9yZy94MjY0Lmh0bWwgLSBvcHRpb25zOiBjYWJhYz0wIHJlZj0zIGRl"
    "YmxvY2s9MTowOjAgYW5hbHlzZT0weDE6MHgxMTEgbWU9aGV4IHN1Ym1lPTcgcHN5PTEgcHN5"
    "X3JkPTEuMDA6MC4wMCBtaXhlZF9yZWY9MSBtZV9yYW5nZT0xNiBjaHJvbWFfbWU9MSB0cmVs"
    "bGlzPTEgOHg4ZGN0PTAgY3FtPTAgZGVhZHpvbmU9MjEsMTEgZmFzdF9wc2tpcD0xIGNocm9t"
    "YV9xcF9vZmZzZXQ9LTIgdGhyZWFkcz00IGxvb2thaGVhZF90aHJlYWRzPTEgc2xpY2VkX3Ro"
    "cmVhZHM9MCBucj0wIGRlY2ltYXRlPTEgaW50ZXJsYWNlZD0wIGJsdXJheV9jb21wYXQ9MCBj"
    "b25zdHJhaW5lZF9pbnRyYT0wIGJmcmFtZXM9MCB3ZWlnaHRwPTAga2V5aW50PTI1MCBrZXlp"
    "bnRfbWluPTggc2NlbmVjdXQ9NDAgaW50cmFfcmVmcmVzaD0wIHJjX2xvb2thaGVhZD00MCBy"
    "Yz1jcmYgbWJ0cmVlPTEgY3JmPTIzLjAgcWNvbXA9MC42MCBxcG1pbj0wIHFwbWF4PTY5IHFw"
    "c3RlcD00IGlwX3JhdGlvPTEuNDAgYXE9MToxLjAwAIAAAABRZYiEBDxGKAAJEscAATgo4AAk"
    "lycnJycnJycnJ11111111111111111111111111111111111111111111111111111111111"
    "111111111114AAAAB0GaOAh4FGAAAAAHQZpUAh4FGAAAAAdBmmAQ8CjAAAAAB0GagBDwKMAA"
    "AAAHQZqgEPAowAAAAAdBmsAQ8CjAAAAAB0Ga4BDwKMAAAAAHQZsAEPAowAAAAAdBmyAQ8CjA"
    "AAAAB0GbQBDwKMAAAAAHQZtgEPAowAAAAAdBm4AQ8CjAAAAABkGboD/AowAAAAZBm8A/wKMA"
    "AAAGQZvgO8Cj"
        ]
    )
)


# Every block the rendered representation knows, in one document: the
# remote image at the end is there to show that it is blocked, not fetched.
SAMPLE_MARKDOWN = r"""# Sample document

This file exists so the rendered Markdown representation has something
to show. Its source is one click away in the editor.

```summary
Every block this viewer knows, in one file, so a change to any of them has
somewhere to be seen.
```

```decision
Decision: Keep documents portable Markdown
Reason: A callout is a fenced code block anywhere else, so nothing is lost.
```

## What it renders

| Block | Comes from |
| --- | --- |
| Diagram | a `mermaid` fence |
| Chart | an `echarts` fence |
| Formula | KaTeX |

- [x] tables, lists and task lists
- [x] fenced code with the editor's own highlighting
- [x] a table of contents, search and export

The mass-energy equivalence is $E = mc^2$, and a block formula reads

$$
\sum_{i=1}^{n} i = \frac{n(n+1)}{2}
$$

## A diagram

```mermaid
graph LR
  A[Catalog] --> B[Document]
  B --> C[Editor]
```

## A chart

```echarts
{
  "xAxis": { "type": "category", "data": ["Q1", "Q2", "Q3"] },
  "yAxis": { "type": "value" },
  "series": [{ "type": "bar", "data": [820, 932, 1200] }]
}
```

## A drawing

```svg
<svg viewBox="0 0 240 80" width="240" height="80" xmlns="http://www.w3.org/2000/svg">
  <circle cx="40" cy="40" r="24" fill="#6ea8fe"/>
  <circle cx="120" cy="40" r="24" fill="#c6f24e"/>
  <circle cx="200" cy="40" r="24" fill="#f78c6c"/>
</svg>
```

## Some code

```python
def total(rows: list[int]) -> int:
    return sum(rows)
```

A [relative link](notes.txt), a local image beside this document, and a remote
image, which is blocked:

![local logo](logo.png)

![remote](https://example.invalid/logo.png)
"""


def write_samples(home: Path) -> None:
    """Write examples for the native viewers and explicit table projections."""
    home.mkdir(mode=0o700, parents=True, exist_ok=True)

    rows = "\n".join(
        f"{index},city-{index},{index * 1000}" for index in range(1, SAMPLE_ROWS + 1)
    )
    (home / "sample.csv").write_text(
        f"id,name,population\n{rows}\n", encoding="utf-8"
    )
    (home / "sample.tsv").write_text(
        "id\tname\n1\tone\n2\ttwo\n", encoding="utf-8"
    )
    # An export from somewhere else, keyed by what the demonstration schema
    # already calls an organization. It is here so that a declared reference
    # (D106) has a pair that actually matches: this file's `slug` and
    # `demo_lab.organizations.slug` are the same string, and neither source
    # knows anything about the other.
    (home / "accounts.csv").write_text(
        "slug,account_manager,renewal_on,seats\n"
        "northstar-labs,Ana Kovalenko,2027-01-31,120\n"
        "blue-orbit,Chen Wei,2026-11-15,45\n",
        encoding="utf-8",
    )
    (home / "notes.txt").write_text(
        "\n".join(f"line {index}" for index in range(1, 21)) + "\n", encoding="utf-8"
    )
    # Structured content with a nested value and a number a float would round.
    # Its alternative table projection also has a key only a later record carries.
    people = [
        {"id": 1, "name": "Ann", "tags": ["reader"]},
        {"id": 2, "name": "Bo", "city": "Riga", "tags": []},
        {"id": 3, "name": "Cyd", "account": 10000000000000001},
    ]
    (home / "people.json").write_text(
        json.dumps(people, indent=2) + "\n", encoding="utf-8"
    )
    events = [
        {"at": f"2026-08-{day:02d}", "kind": "login", "ok": day % 2 == 0}
        for day in range(1, 6)
    ]
    (home / "events.jsonl").write_text(
        "".join(f"{json.dumps(event)}\n" for event in events), encoding="utf-8"
    )
    # Named .json and not JSON: its JSON viewer shows the source and says why.
    (home / "broken.json").write_text('{"a": 1,\n "b" 2}\n', encoding="utf-8")
    # A log, which reads as text and opens as a table of the columns its own
    # lines are written in. Two continuation lines, because a real log has
    # them and they are what a naive parser drops.
    (home / "server.log").write_text(
        "".join(
            f"2026-08-{day:02d} 09:1{day}:00.100 UTC [{100 + day}] "
            f"{level}:  {message}\n" + detail
            for day, level, message, detail in SAMPLE_LOG_LINES
        ),
        encoding="utf-8",
    )
    (home / "logo.png").write_bytes(SAMPLE_PNG)
    (home / "guide.md").write_text(SAMPLE_MARKDOWN, encoding="utf-8")

    _write_dataset_samples(home)
    _write_media_samples(home)
    _write_archive_sample(home)

    reports = home / "reports"
    reports.mkdir(exist_ok=True)
    (reports / "q3.csv").write_text("quarter,revenue\nQ3,1200\n", encoding="utf-8")
    (home / "empty").mkdir(exist_ok=True)


def _write_dataset_samples(home: Path) -> None:
    """A columnar file and a workbook, for the datasets of D51.

    Written here rather than committed as fixtures so the samples stay honest
    about the libraries that read them, and so nothing binary lives in the
    repository. This is the one place outside the file adapter that touches
    either library, and it is a development tool rather than the application.
    """
    import pyarrow as pa
    import pyarrow.parquet as pq
    from openpyxl import Workbook

    table = pa.table(
        {
            "id": pa.array(range(1, SAMPLE_ROWS + 1), type=pa.int64()),
            "city": pa.array([f"city-{index}" for index in range(1, SAMPLE_ROWS + 1)]),
            "population": pa.array(
                [index * 1000.5 for index in range(1, SAMPLE_ROWS + 1)],
                type=pa.float64(),
            ),
            "capital": pa.array([index % 7 == 0 for index in range(1, SAMPLE_ROWS + 1)]),
            "tags": pa.array(
                [["north"] if index % 2 else ["south", "coast"] for index in range(1, SAMPLE_ROWS + 1)],
                type=pa.list_(pa.string()),
            ),
        }
    )
    # Several row groups on purpose: a page then begins inside one of them.
    pq.write_table(table, home / "cities.parquet", row_group_size=100)

    book = Workbook()
    revenue = book.active
    if revenue is None:  # pragma: no cover - a new workbook always has one sheet
        raise RuntimeError("A new workbook has no active sheet")
    revenue.title = "Revenue"
    revenue.append(["quarter", "amount", "closed"])
    for quarter, amount in (("Q1", 820), ("Q2", 932), ("Q3", 1200), ("Q4", 1310)):
        revenue.append([quarter, amount, quarter != "Q4"])
    people = book.create_sheet("People")
    people.append(["name", "started"])
    people.append(["Ann", datetime(2026, 1, 14, tzinfo=UTC).date()])
    people.append(["Bo", datetime(2026, 3, 2, tzinfo=UTC).date()])
    hidden = book.create_sheet("Working notes")
    hidden.append(["note"])
    hidden.append(["not for the report"])
    hidden.sheet_state = "hidden"
    book.save(home / "budget.xlsx")


def _write_media_samples(home: Path) -> None:
    """A sound file, a recording and a document, for the viewers of D52."""
    with wave.open(str(home / "tone.wav"), "wb") as sound:
        sound.setnchannels(1)
        sound.setsampwidth(2)
        sound.setframerate(SAMPLE_RATE)
        sound.writeframes(_tone(seconds=2, hertz=440))
    (home / "clip.mp4").write_bytes(SAMPLE_MP4)
    (home / "guide.pdf").write_bytes(_pdf(SAMPLE_PDF_PAGES))


def _tone(*, seconds: int, hertz: int) -> bytes:
    """A sine wave, so the sample is something a person can hear."""
    frames = bytearray()
    for index in range(seconds * SAMPLE_RATE):
        value = int(12000 * math.sin(2 * math.pi * hertz * index / SAMPLE_RATE))
        frames += struct.pack("<h", value)
    return bytes(frames)


def _pdf(pages: list[list[str]]) -> bytes:
    """A real PDF of a few lines a page, written out object by object.

    Small enough to keep here and honest enough for the viewer to draw: the
    cross-reference table is computed from the offsets, as a reader expects.
    """
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # the page tree, once the page objects are numbered
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids: list[str] = []
    for lines in pages:
        content = "BT /F1 16 Tf 72 760 Td 20 TL\n"
        content += "".join(f"({line}) Tj T*\n" for line in lines)
        content += "ET"
        stream = content.encode("ascii")
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        )
        page_number = len(objects) + 1
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents "
            + str(page_number - 1).encode()
            + b" 0 R >>"
        )
        kids.append(f"{page_number} 0 R")
    objects[1] = (
        b"<< /Type /Pages /Kids [" + " ".join(kids).encode() + b"] /Count "
        + str(len(pages)).encode()
        + b" >>"
    )

    out = bytearray(b"%PDF-1.7\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += str(number).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    start = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n"
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size " + str(len(objects) + 1).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(start).encode()
        + b"\n%%EOF\n"
    )
    return bytes(out)


def _write_archive_sample(home: Path) -> None:
    """One archive with a folder in it, for the browsing of D53."""
    with zipfile.ZipFile(home / "bundle.zip", "w", zipfile.ZIP_DEFLATED) as bundle:
        bundle.writestr("readme.md", SAMPLE_ARCHIVE_README)
        bundle.writestr("docs/notes.txt", "one\ntwo\nthree\n")
        bundle.writestr("docs/cities.csv", "id,name\n1,one\n2,two\n3,three\n")
        bundle.writestr("docs/logo.png", SAMPLE_PNG)


def remove_orphan_homes(*, keep: set[uuid.UUID]) -> list[uuid.UUID]:
    """Delete the home directories of users that no longer exist.

    The test suite shares the development database and empties the user table,
    so every run strands a full directory under an id nobody can log in as.

    This decides *which* homes go; `storage.remove_home` is what takes one
    away, so where a home lives and how it is removed keep one definition. Only
    a real directory named after a user id is a candidate: anything else under
    the root was put there for a reason this script does not know, and guessing
    about it is the one destructive mistake it could make.
    """
    root = settings.USER_STORAGE_ROOT
    if not root.is_dir():
        return []

    removed: list[uuid.UUID] = []
    for entry in sorted(root.iterdir()):
        if entry.is_symlink() or not entry.is_dir():
            logger.info("Not a home directory, left alone: %s", entry.name)
            continue
        try:
            user_id = uuid.UUID(entry.name)
        except ValueError:
            logger.info("Name is not a user id, left alone: %s", entry.name)
            continue
        if user_id in keep:
            continue
        storage.remove_home(user_id)
        removed.append(user_id)
        logger.info("Removed the home of a user that no longer exists: %s", entry.name)
    return removed


async def main(email: str) -> None:
    async with async_session_maker() as session:
        user = await crud.get_user_by_email(session=session, email=email)
        live_ids = await crud.get_all_user_ids(session=session)
    if user is None:
        logger.error("No such user: %s", email)
        raise SystemExit(1)

    removed = remove_orphan_homes(keep=live_ids)
    logger.info("Removed %d orphaned home directories", len(removed))

    home = storage.home_dir(user.id)
    write_samples(home)
    logger.info("Sample files written to %s", home)


if __name__ == "__main__":
    address = sys.argv[1] if len(sys.argv) > 1 else settings.FIRST_SUPERUSER
    asyncio.run(main(address), loop_factory=loop_factory)
